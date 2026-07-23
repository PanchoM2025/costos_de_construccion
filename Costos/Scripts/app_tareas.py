import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Data", "Master_FINAL.db")
CODIGOS_GENERICOS = {"CH", "MOP"}


def get_conn():
    conn = sqlite3.connect(DB_PATH)
    return conn


def cod_tar_existe(cod):
    with get_conn() as conn:
        row = conn.execute("SELECT 1 FROM tareas WHERE cod_tar=? LIMIT 1", (cod,)).fetchone()
    return row is not None


def _ensure_extra_columns():
    with get_conn() as conn:
        cols = [r[1] for r in conn.execute("PRAGMA table_info(tareas)").fetchall()]
        if "cantidad" not in cols:
            try:
                conn.execute("ALTER TABLE tareas ADD COLUMN cantidad REAL")
                conn.commit()
            except Exception:
                pass


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Gestión de Tareas e Insumos")
        self.geometry("1220x680"); self.minsize(900, 520); self.configure(bg="#F0F2F5")
        self._current_cod_tar = None; self._current_tar_id = None
        self._apply_styles(); self._build_ui(); _ensure_extra_columns(); self._load_tareas()

    def _apply_styles(self):
        s = ttk.Style(self); s.theme_use("clam")
        s.configure("Treeview", rowheight=24, font=("Segoe UI",10), fieldbackground="#FFFFFF")
        s.configure("Treeview.Heading", font=("Segoe UI",10,"bold"), background="#2C3E50", foreground="white", relief="flat")
        s.map("Treeview", background=[("selected","#2980B9")], foreground=[("selected","white")])
        s.configure("TLabel", background="#F0F2F5", font=("Segoe UI",10))
        s.configure("TEntry", font=("Segoe UI",10))
        s.configure("TButton", font=("Segoe UI",10), padding=(8,5))
        s.configure("TLabelframe", background="#F0F2F5")
        s.configure("TLabelframe.Label", font=("Segoe UI",10,"bold"), background="#F0F2F5", foreground="#2C3E50")
        s.configure("Add.TButton", background="#27AE60", foreground="white", font=("Segoe UI",10,"bold"))
        s.map("Add.TButton", background=[("active","#219A52"),("disabled","#BDC3C7")], foreground=[("disabled","#888")])
        s.configure("Edit.TButton", background="#2980B9", foreground="white", font=("Segoe UI",10,"bold"))
        s.map("Edit.TButton", background=[("active","#2471A3"),("disabled","#BDC3C7")], foreground=[("disabled","#888")])
        s.configure("Del.TButton", background="#E74C3C", foreground="white", font=("Segoe UI",10,"bold"))
        s.map("Del.TButton", background=[("active","#CB4335"),("disabled","#BDC3C7")], foreground=[("disabled","#888")])
        s.configure("Export.TButton", background="#8E44AD", foreground="white", font=("Segoe UI",10,"bold"))
        s.map("Export.TButton", background=[("active","#7D3C98"),("disabled","#BDC3C7")], foreground=[("disabled","#888")])

    def _build_ui(self):
        hdr = tk.Frame(self, bg="#2C3E50", height=50); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr, text="  Gestión de Tareas e Insumos", bg="#2C3E50", fg="white",
                 font=("Segoe UI",14,"bold")).pack(side=tk.LEFT, padx=8, pady=10)
        body = tk.Frame(self, bg="#F0F2F5"); body.pack(fill=tk.BOTH, expand=True, padx=12, pady=10)
        lf_tar = ttk.LabelFrame(body, text="Tareas", padding=8)
        lf_tar.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0,8))
        sf = tk.Frame(lf_tar, bg="#F0F2F5"); sf.pack(fill=tk.X, pady=(0,6))
        ttk.Label(sf, text="Buscar:").pack(side=tk.LEFT)
        self.sv_search = tk.StringVar()
        self.sv_search.trace_add("write", lambda *_: self._filter_tareas())
        ttk.Entry(sf, textvariable=self.sv_search).pack(side=tk.LEFT, padx=6, fill=tk.X, expand=True)
        frame_tar = tk.Frame(lf_tar, bg="#F0F2F5"); frame_tar.pack(fill=tk.BOTH, expand=True)
        self.tv_tar = ttk.Treeview(frame_tar, columns=("cod","desc","unidad"), show="headings", selectmode="extended")
        self.tv_tar.heading("cod", text="Código"); self.tv_tar.heading("desc", text="Descripción"); self.tv_tar.heading("unidad", text="Unidad")
        self.tv_tar.column("cod", width=80, anchor=tk.CENTER, stretch=False)
        self.tv_tar.column("desc", width=310, minwidth=140); self.tv_tar.column("unidad", width=65, anchor=tk.CENTER, stretch=False)
        sb_t = ttk.Scrollbar(frame_tar, orient=tk.VERTICAL, command=self.tv_tar.yview)
        self.tv_tar.configure(yscrollcommand=sb_t.set)
        self.tv_tar.pack(side=tk.LEFT, fill=tk.BOTH, expand=True); sb_t.pack(side=tk.LEFT, fill=tk.Y)
        self.tv_tar.bind("<<TreeviewSelect>>", self._on_tar_select)
        self.tv_tar.bind("<Button-3>", self._on_tar_rightclick)
        btn_frm_tar = tk.Frame(lf_tar, bg="#F0F2F5"); btn_frm_tar.pack(fill=tk.X, pady=(8,0))
        self.btn_modificar_tar = ttk.Button(btn_frm_tar, text="✏ Modificar", style="Edit.TButton", state=tk.DISABLED, command=self._open_modificar_tar)
        self.btn_modificar_tar.pack(side=tk.LEFT, padx=(0,4))
        self.btn_agregar_tar = ttk.Button(btn_frm_tar, text="➕ Agregar", style="Edit.TButton", command=self._open_agregar_tar)
        self.btn_agregar_tar.pack(side=tk.LEFT)
        self.btn_exportar = ttk.Button(btn_frm_tar, text="📤 Exportar", style="Export.TButton", state=tk.DISABLED, command=self._export_tareas)
        self.btn_exportar.pack(side=tk.LEFT, padx=(8,0))
        lf_ins = ttk.LabelFrame(body, text="Insumos de la tarea", padding=8); lf_ins.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.lbl_tar = ttk.Label(lf_ins, text="← Seleccioná una tarea", font=("Segoe UI",10,"italic"), foreground="#7F8C8D")
        self.lbl_tar.pack(anchor=tk.W, pady=(0,6))
        frame_ins = tk.Frame(lf_ins, bg="#F0F2F5"); frame_ins.pack(fill=tk.BOTH, expand=True)
        self.tv_ins = ttk.Treeview(frame_ins, columns=("cod_ins","desc_ins","unidad_ins","cantidad"), show="headings", selectmode="browse")
        for col,w,txt,anc in [("cod_ins",80,"Código",tk.CENTER),("desc_ins",230,"Descripción",tk.W),("unidad_ins",65,"Unidad",tk.CENTER),("cantidad",80,"Cantidad",tk.CENTER)]:
            self.tv_ins.heading(col, text=txt); self.tv_ins.column(col, width=w, anchor=anc, stretch=(col=="desc_ins"))
        sb_i = ttk.Scrollbar(frame_ins, orient=tk.VERTICAL, command=self.tv_ins.yview)
        self.tv_ins.configure(yscrollcommand=sb_i.set)
        self.tv_ins.pack(side=tk.LEFT, fill=tk.BOTH, expand=True); sb_i.pack(side=tk.LEFT, fill=tk.Y)
        self.tv_ins.bind("<<TreeviewSelect>>", self._on_ins_select)
        btn_frm = tk.Frame(lf_ins, bg="#F0F2F5"); btn_frm.pack(fill=tk.X, pady=(8,0))
        self.btn_add = ttk.Button(btn_frm, text="＋  Agregar", style="Add.TButton", state=tk.DISABLED, command=self._open_agregar)
        self.btn_add.pack(side=tk.LEFT, padx=(0,4))
        self.btn_edit = ttk.Button(btn_frm, text="✎  Editar", style="Edit.TButton", state=tk.DISABLED, command=self._open_editar)
        self.btn_edit.pack(side=tk.LEFT, padx=(0,4))
        self.btn_del = ttk.Button(btn_frm, text="✕  Eliminar", style="Del.TButton", state=tk.DISABLED, command=self._eliminar_insumo)
        self.btn_del.pack(side=tk.LEFT, padx=(0,4))
        self.btn_rename = ttk.Button(btn_frm, text="Renombrar", style="Edit.TButton", state=tk.DISABLED, command=self._open_renombrar_main)
        self.btn_rename.pack(side=tk.LEFT, padx=(0,4))
        self.btn_nuevo_cod_ins = ttk.Button(btn_frm, text="Nuevo COD_INS", style="Edit.TButton", state=tk.DISABLED, command=self._open_nuevo_cod_ins)
        self.btn_nuevo_cod_ins.pack(side=tk.LEFT)
        self.sv_status = tk.StringVar(value="Listo")
        tk.Label(self, textvariable=self.sv_status, anchor=tk.W, bg="#2C3E50", fg="white",
                 font=("Segoe UI",9), padx=8).pack(fill=tk.X, side=tk.BOTTOM)

    def _load_tareas(self):
        with get_conn() as conn:
            tiene_rubros = conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='rubros'").fetchone()[0]
            if tiene_rubros:
                rows = conn.execute("""
                    SELECT t.id, t.cod_tar, t.descripcion, t.unidad,
                           COALESCE(rp.designacion, r.designacion, '') AS rubro_desc,
                           COALESCE(r.cod_rubro, '') AS sub_cod
                    FROM tareas t
                    LEFT JOIN rubros r ON r.id_rs_ini = t.id_rs
                        AND (r.es_padre=0 OR NOT EXISTS (SELECT 1 FROM rubros r2 WHERE r2.id_rs_ini=t.id_rs AND r2.es_padre=0))
                    LEFT JOIN rubros rp ON rp.cod_rubro = r.cod_padre
                    WHERE t.activo=1
                    ORDER BY COALESCE(rp.id_rs_ini,r.id_rs_ini,'ZZ'), COALESCE(r.cod_rubro,''),
                             t.id_tarea NULLS LAST, t.cod_tar, t.id
                """).fetchall()
            else:
                rows = [(r[0],r[1],r[2],r[3],'','') for r in conn.execute(
                    "SELECT id,cod_tar,descripcion,unidad FROM tareas WHERE activo=1 ORDER BY id").fetchall()]
        self._all_tareas = rows; self._populate_tareas(rows)

    def _populate_tareas(self, rows):
        self.tv_tar.delete(*self.tv_tar.get_children())
        self.tv_tar.tag_configure("generico", foreground="#E67E22")
        self.tv_tar.tag_configure("grupo_hdr", background="#1F5C99", foreground="#FFFFFF", font=("Segoe UI",9,"bold"))
        grupo_actual = None; grupo_seq = 0; tar_count = 0
        for row in rows:
            tar_id, cod, desc, unidad, rubro_desc, sub_cod = row
            grupo = f"{sub_cod}  —  {rubro_desc}" if sub_cod else (rubro_desc or "")
            if grupo != grupo_actual:
                grupo_actual = grupo; grupo_seq += 1
                self.tv_tar.insert("", tk.END, iid=f"_grp_{grupo_seq}", values=("",grupo,""), tags=("grupo_hdr",))
            tag = "generico" if cod in CODIGOS_GENERICOS else ""
            self.tv_tar.insert("", tk.END, iid=str(tar_id), values=(cod, desc or "", unidad or ""), tags=(tag,))
            tar_count += 1
        self.sv_status.set(f"{tar_count} tareas")

    def _filter_tareas(self):
        q = self.sv_search.get().lower()
        filtered = self._all_tareas if not q else [
            r for r in self._all_tareas if q in r[1].lower() or q in (r[2] or "").lower() or q in (r[4] or "").lower()]
        self._populate_tareas(filtered)

    def _on_tar_select(self, _=None):
        sel = self.tv_tar.selection()
        if not sel: return
        grp_sels = [s for s in sel if str(s).startswith("_grp_")]
        if grp_sels: self.tv_tar.selection_remove(*grp_sels); sel = [s for s in sel if not str(s).startswith("_grp_")]
        if not sel: return
        self.btn_exportar.config(state=tk.NORMAL)
        if len(sel) == 1:
            tar_id = int(sel[0]); vals = self.tv_tar.item(sel[0], "values")
            self._current_tar_id = tar_id; self._current_cod_tar = vals[0]
            self.lbl_tar.config(text=f"{vals[0]}  —  {vals[1]}", foreground="#E67E22" if vals[0] in CODIGOS_GENERICOS else "#2C3E50")
            self.btn_add.config(state=tk.NORMAL); self.btn_edit.config(state=tk.DISABLED)
            self.btn_del.config(state=tk.DISABLED); self.btn_modificar_tar.config(state=tk.NORMAL)
            self._load_insumos(vals[0])
        else:
            self.btn_modificar_tar.config(state=tk.DISABLED); self.btn_add.config(state=tk.DISABLED)
            self.lbl_tar.config(text=f"{len(sel)} tareas seleccionadas", foreground="#2C3E50")
            self.sv_status.set(f"{len(sel)} tareas seleccionadas — hacé clic en Exportar")

    def _on_tar_rightclick(self, event):
        item = self.tv_tar.identify_row(event.y)
        if not item or str(item).startswith("_grp_"): return
        vals = self.tv_tar.item(item, "values")
        if not vals or not vals[0] or vals[0] in CODIGOS_GENERICOS: return
        self.tv_tar.selection_set(item); self._on_tar_select()
        InsumoContextDialog(self, tar_id=int(item), cod_tar=vals[0], desc_tar=vals[1], on_incorporar=self._open_incorporar_tar)

    def _open_incorporar_tar(self, tar_id):
        with get_conn() as conn:
            row = conn.execute("SELECT id,cod_tar,descripcion,unidad,id_rs FROM tareas WHERE id=?", (tar_id,)).fetchone()
        if not row: return
        _, cod_tar, desc, unid, id_rs = row
        ModificarTareaDialog(self, tar_id=tar_id, row=(cod_tar,desc,unid,id_rs), on_guardado=self._on_tarea_modificada)

    def _export_tareas(self):
        sel = [s for s in self.tv_tar.selection() if not str(s).startswith("_grp_")]
        if not sel: messagebox.showwarning("Exportar","Seleccioná al menos una tarea.",parent=self); return
        filepath = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")],
            initialfile="tareas_exportadas.xlsx", parent=self, title="Guardar exportación de tareas")
        if not filepath: return
        tar_ids = [int(s) for s in sel]; placeholders = ",".join("?"*len(tar_ids))
        try:
            with get_conn() as conn:
                tiene_rubros = conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='rubros'").fetchone()[0]
                query = f"""
                    SELECT COALESCE(rp.cod_rubro,r.cod_rubro,'') AS cod_rubro,
                           COALESCE(rp.designacion,r.designacion,'') AS desc_rubro,
                           CASE WHEN rp.cod_rubro IS NOT NULL THEN r.cod_rubro ELSE '' END AS cod_sub,
                           CASE WHEN rp.cod_rubro IS NOT NULL THEN r.designacion ELSE '' END AS desc_sub,
                           t.cod_tar, t.descripcion, t.unidad,
                           inc.cod_ins, COALESCE(ins.descripcion,''), COALESCE(ins.unidad,inc.unidad,''), inc.cantidad
                    FROM tareas t
                    LEFT JOIN rubros r ON r.id_rs_ini=t.id_rs AND (r.es_padre=0 OR NOT EXISTS (SELECT 1 FROM rubros r2 WHERE r2.id_rs_ini=t.id_rs AND r2.es_padre=0))
                    LEFT JOIN rubros rp ON rp.cod_rubro=r.cod_padre
                    LEFT JOIN incidencias inc ON inc.cod_tar=t.cod_tar
                    LEFT JOIN insumos ins ON ins.cod_ins=inc.cod_ins
                    WHERE t.id IN ({placeholders})
                    ORDER BY COALESCE(rp.id_rs_ini,r.id_rs_ini,'ZZ'), COALESCE(r.cod_rubro,''), t.id_tarea NULLS LAST, t.cod_tar, inc.cod_ins
                """ if tiene_rubros else f"""
                    SELECT '','','','', t.cod_tar, t.descripcion, t.unidad,
                           inc.cod_ins, COALESCE(ins.descripcion,''), COALESCE(ins.unidad,inc.unidad,''), inc.cantidad
                    FROM tareas t LEFT JOIN incidencias inc ON inc.cod_tar=t.cod_tar
                    LEFT JOIN insumos ins ON ins.cod_ins=inc.cod_ins
                    WHERE t.id IN ({placeholders}) ORDER BY t.cod_tar, inc.cod_ins"""
                rows = conn.execute(query, tar_ids).fetchall()
            C_HDR='0D3D6B'; C_RUBRO='1F5C99'; C_SUB='2E86C1'; C_TAR='D6E8F5'; C_ODD='FFFFFF'; C_EVEN='F5F9FF'
            def fill(c): return PatternFill(fill_type="solid", fgColor=c)
            wb = openpyxl.Workbook(); ws = wb.active; ws.title="Tareas"
            headers=["Rubro","Desc. Rubro","SubRubro","Desc. SubRubro","COD_TAR","Desc. Tarea","Unidad","COD_INS","Desc. Insumo","Unidad Ins","Cantidad"]
            col_widths=[10,42,10,42,12,52,9,12,52,11,11]; letters=list("ABCDEFGHIJK")
            hdr_font=Font(bold=True,color="FFFFFF",name="Calibri",size=11)
            rub_font=Font(bold=True,color="FFFFFF",name="Calibri",size=11)
            sub_font=Font(bold=True,color="FFFFFF",name="Calibri",size=10)
            tar_font=Font(bold=True,color="2C3E50",name="Calibri",size=10)
            ins_font=Font(name="Calibri",size=10)
            ws.row_dimensions[1].height=22
            for col,(h,w,ltr) in enumerate(zip(headers,col_widths,letters),1):
                cell=ws.cell(row=1,column=col,value=h); cell.font=hdr_font; cell.fill=fill(C_HDR)
                cell.alignment=Alignment(horizontal="center",vertical="center"); ws.column_dimensions[ltr].width=w
            row_idx=2; last_rubro=object(); last_sub=object(); last_tar=object(); ins_parity=0
            for r in rows:
                cod_rub,desc_rub,cod_sub,desc_sub,cod_tar,desc_tar,unid_tar,cod_ins,desc_ins,unid_ins,cantidad = r
                if (cod_rub,desc_rub) != last_rubro:
                    last_rubro=(cod_rub,desc_rub); last_sub=object(); last_tar=object(); ins_parity=0
                    f=fill(C_RUBRO); ws.cell(row=row_idx,column=1,value=cod_rub or "").font=rub_font
                    ws.cell(row=row_idx,column=2,value=desc_rub or "").font=rub_font
                    for c in range(1,12): ws.cell(row=row_idx,column=c).fill=f
                    ws.row_dimensions[row_idx].height=20; row_idx+=1
                if (cod_sub,desc_sub) != last_sub:
                    last_sub=(cod_sub,desc_sub); last_tar=object(); ins_parity=0
                    f=fill(C_SUB); ws.cell(row=row_idx,column=3,value=cod_sub or "").font=sub_font
                    ws.cell(row=row_idx,column=4,value=desc_sub or "").font=sub_font
                    for c in range(1,12): ws.cell(row=row_idx,column=c).fill=f
                    ws.row_dimensions[row_idx].height=18; row_idx+=1
                if cod_tar != last_tar:
                    last_tar=cod_tar; ins_parity=0; f=fill(C_TAR)
                    ws.cell(row=row_idx,column=5,value=cod_tar or "").font=tar_font
                    ws.cell(row=row_idx,column=6,value=desc_tar or "").font=tar_font
                    ws.cell(row=row_idx,column=7,value=unid_tar or "").font=tar_font
                    for c in range(1,12): ws.cell(row=row_idx,column=c).fill=f
                    ws.row_dimensions[row_idx].height=18; row_idx+=1
                if cod_ins:
                    f=fill(C_ODD if ins_parity%2==0 else C_EVEN)
                    ws.cell(row=row_idx,column=8,value=cod_ins or "").font=ins_font
                    ws.cell(row=row_idx,column=9,value=desc_ins or "").font=ins_font
                    ws.cell(row=row_idx,column=10,value=unid_ins or "").font=ins_font
                    qty_c=ws.cell(row=row_idx,column=11,value=float(cantidad) if cantidad is not None else "")
                    qty_c.font=ins_font
                    if cantidad is not None: qty_c.number_format="0.0000"; qty_c.alignment=Alignment(horizontal="right")
                    for c in range(1,12): ws.cell(row=row_idx,column=c).fill=f
                    row_idx+=1; ins_parity+=1
            ws.freeze_panes="A2"; wb.save(filepath)
            n_t=len({r[4] for r in rows})
            messagebox.showinfo("Exportar",f"{n_t} tarea(s) exportadas.\n\n{filepath}",parent=self)
            self.sv_status.set(f"✓ {n_t} tareas exportadas → {os.path.basename(filepath)}")
        except Exception as e:
            messagebox.showerror("Error al exportar",str(e),parent=self)

    def _load_insumos(self, cod_tar):
        self.tv_ins.delete(*self.tv_ins.get_children())
        with get_conn() as conn:
            rows = conn.execute("""
                SELECT inc.cod_ins, COALESCE(ins.descripcion,'(sin descripción)'), inc.unidad, inc.cantidad
                FROM incidencias inc LEFT JOIN insumos ins ON ins.cod_ins=inc.cod_ins
                WHERE inc.cod_tar=? ORDER BY inc.cod_ins""", (cod_tar,)).fetchall()
        for i,(cod_ins,desc_ins,unidad,cantidad) in enumerate(rows):
            self.tv_ins.insert("",tk.END,iid=f"i_{i}",values=(cod_ins,desc_ins,unidad or "",f"{cantidad:.4f}" if cantidad is not None else ""))
        n=len(rows); self.sv_status.set(f"{n} insumo{'s' if n!=1 else ''}  ·  tarea {cod_tar}")

    def _on_ins_select(self, _=None):
        ok = bool(self.tv_ins.selection() and self._current_cod_tar)
        state = tk.NORMAL if ok else tk.DISABLED
        self.btn_edit.config(state=state); self.btn_del.config(state=state)
        self.btn_rename.config(state=state); self.btn_nuevo_cod_ins.config(state=state)

    def _get_selected_ins(self):
        sel = self.tv_ins.selection()
        return self.tv_ins.item(sel[0],"values") if sel else None

    def _eliminar_insumo(self):
        vals = self._get_selected_ins()
        if not vals: return
        cod_ins = vals[0]
        if not messagebox.askyesno("Confirmar eliminación",f"¿Eliminás el insumo  {cod_ins}  de la tarea  {self._current_cod_tar}?",parent=self): return
        with get_conn() as conn:
            conn.execute("DELETE FROM incidencias WHERE cod_tar=? AND cod_ins=?",(self._current_cod_tar,cod_ins))
        self._load_insumos(self._current_cod_tar)
        self.btn_edit.config(state=tk.DISABLED); self.btn_del.config(state=tk.DISABLED)
        self.sv_status.set(f"Insumo {cod_ins} eliminado de {self._current_cod_tar}.")

    def _open_modificar_tar(self):
        if not self._current_cod_tar: return
        with get_conn() as conn:
            row = conn.execute("SELECT id,cod_tar,descripcion,unidad,id_rs FROM tareas WHERE cod_tar=?",(self._current_cod_tar,)).fetchone()
        if not row: return
        tar_id,cod_tar,desc,unid,id_rs = row
        ModificarTareaDialog(self,tar_id=tar_id,row=(cod_tar,desc,unid,id_rs),on_guardado=self._on_tarea_modificada)

    def _on_tarea_modificada(self,cod_viejo,cod_nuevo):
        self._current_cod_tar=cod_nuevo; self._load_tareas()
        for item in self.tv_tar.get_children():
            vals=self.tv_tar.item(item,"values")
            if vals and vals[0]==cod_nuevo: self.tv_tar.selection_set(item); self.tv_tar.see(item); break
        self._load_insumos(cod_nuevo); self.sv_status.set(f"Tarea modificada: {cod_viejo} → {cod_nuevo}")

    def _open_agregar_tar(self): AgregarTareaDialog(self, on_guardado=self._on_tarea_agregada)

    def _on_tarea_agregada(self, cod_nuevo):
        self._load_tareas()
        for item in self.tv_tar.get_children():
            vals=self.tv_tar.item(item,"values")
            if vals and vals[0]==cod_nuevo: self.tv_tar.selection_set(item); self.tv_tar.see(item); break
        self.sv_status.set(f"Tarea agregada: {cod_nuevo}")

    def _open_nuevo_cod_tar(self): self._open_modificar_tar()
    def _on_cod_tar_cambiado(self,v,n): self._on_tarea_modificada(v,n)

    def _open_nuevo_cod_ins(self):
        vals=self._get_selected_ins()
        if not vals: return
        CambiarCodigoDialog(self,tipo="INS",cod_actual=vals[0],on_cambiado=self._on_cod_ins_cambiado)

    def _on_cod_ins_cambiado(self,cod_viejo,cod_nuevo):
        self._load_insumos(self._current_cod_tar); self.sv_status.set(f"COD_INS cambiado: {cod_viejo} → {cod_nuevo}")

    def _open_renombrar_main(self):
        vals=self._get_selected_ins()
        if not vals: return
        RenombrarInsumoDialog(self,cod_ins=vals[0],desc_actual=vals[1],
                              on_renombrado=lambda c,d: self._load_insumos(self._current_cod_tar))

    def _open_agregar(self):
        if not self._current_cod_tar: return
        if self._current_cod_tar in CODIGOS_GENERICOS:
            AsignarCodigoDialog(self,tar_id=self._current_tar_id,cod_tar_actual=self._current_cod_tar,
                                desc_tar=self.lbl_tar.cget("text").split("  —  ",1)[-1],on_asignado=self._on_codigo_asignado)
        else:
            InsumoDialog(self,mode="add",cod_tar=self._current_cod_tar,on_save=lambda: self._load_insumos(self._current_cod_tar))

    def _on_codigo_asignado(self,nuevo_cod):
        self._current_cod_tar=nuevo_cod
        old_text=self.lbl_tar.cget("text"); desc=old_text.split("  —  ",1)[-1] if "  —  " in old_text else old_text
        self.lbl_tar.config(text=f"{nuevo_cod}  —  {desc}",foreground="#2C3E50"); self._load_tareas()
        InsumoDialog(self,mode="add",cod_tar=nuevo_cod,on_save=lambda: self._load_insumos(nuevo_cod))

    def _open_editar(self):
        vals=self._get_selected_ins()
        if not vals: return
        InsumoDialog(self,mode="edit",cod_tar=self._current_cod_tar,cod_ins=vals[0],cantidad_actual=vals[3],
                     on_save=lambda: self._load_insumos(self._current_cod_tar))


# ═══════════════════════════════════════════════════════════
class InsumoContextDialog(tk.Toplevel):
    def __init__(self,parent,tar_id,cod_tar,desc_tar,on_incorporar):
        super().__init__(parent)
        self.tar_id=tar_id; self.on_incorporar=on_incorporar
        self.title(f"Insumos — {cod_tar}"); self.geometry("640x440"); self.minsize(480,320)
        self.configure(bg="#F0F2F5"); self.grab_set()
        hdr=tk.Frame(self,bg="#1F5C99",height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr,text=f"  {cod_tar}  —  {desc_tar[:55]}",bg="#1F5C99",fg="white",
                 font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        frame_ins=tk.Frame(self,bg="#F0F2F5"); frame_ins.pack(fill=tk.BOTH,expand=True,padx=12,pady=8)
        self.tv=ttk.Treeview(frame_ins,columns=("cod_ins","desc_ins","unidad_ins","cantidad"),show="headings",selectmode="browse")
        for col,w,txt,anc in [("cod_ins",85,"Código",tk.CENTER),("desc_ins",295,"Descripción",tk.W),("unidad_ins",65,"Unidad",tk.CENTER),("cantidad",80,"Cantidad",tk.CENTER)]:
            self.tv.heading(col,text=txt); self.tv.column(col,width=w,anchor=anc,stretch=(col=="desc_ins"))
        sb=ttk.Scrollbar(frame_ins,orient=tk.VERTICAL,command=self.tv.yview)
        self.tv.configure(yscrollcommand=sb.set); self.tv.pack(side=tk.LEFT,fill=tk.BOTH,expand=True); sb.pack(side=tk.LEFT,fill=tk.Y)
        btn_frm=tk.Frame(self,bg="#F0F2F5"); btn_frm.pack(fill=tk.X,padx=12,pady=(0,12))
        ttk.Button(btn_frm,text="Cerrar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(btn_frm,text="  ✏  Incorporar  ",style="Edit.TButton",command=self._incorporar).pack(side=tk.RIGHT)
        with get_conn() as conn:
            rows=conn.execute("SELECT inc.cod_ins,COALESCE(ins.descripcion,'(sin descripción)'),inc.unidad,inc.cantidad FROM incidencias inc LEFT JOIN insumos ins ON ins.cod_ins=inc.cod_ins WHERE inc.cod_tar=? ORDER BY inc.cod_ins",(cod_tar,)).fetchall()
        for i,(c,d,u,q) in enumerate(rows):
            self.tv.insert("",tk.END,iid=f"i_{i}",values=(c,d,u or "",f"{q:.4f}" if q is not None else ""))
        if not rows: self.tv.insert("",tk.END,values=("","(sin insumos)","",""))

    def _incorporar(self):
        cb=self.on_incorporar; tid=self.tar_id; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(tid))


# ═══════════════════════════════════════════════════════════
class AsignarCodigoDialog(tk.Toplevel):
    def __init__(self,parent,tar_id,cod_tar_actual,desc_tar,on_asignado):
        super().__init__(parent)
        self.tar_id=tar_id; self.cod_tar_actual=cod_tar_actual; self.on_asignado=on_asignado
        self.title("Asignar código"); self.geometry("460x220"); self.resizable(False,False)
        self.configure(bg="#F0F2F5"); self.grab_set()
        hdr=tk.Frame(self,bg="#E67E22",height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr,text=f"  Asignar código  ·  {cod_tar_actual} — {desc_tar[:45]}",bg="#E67E22",fg="white",
                 font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        tk.Label(self,text=f"Esta tarea tiene código genérico  \"{cod_tar_actual}\".\nAsignale el código definitivo antes de guardar insumos.",
                 bg="#F0F2F5",fg="#555",font=("Segoe UI",10),justify=tk.LEFT).pack(padx=20,pady=(16,8),anchor=tk.W)
        ef=tk.Frame(self,bg="#F0F2F5"); ef.pack(fill=tk.X,padx=20,pady=(0,6))
        ttk.Label(ef,text="Código nuevo:").pack(side=tk.LEFT)
        self.sv_cod=tk.StringVar()
        e=ttk.Entry(ef,textvariable=self.sv_cod,width=18,font=("Segoe UI",11))
        e.pack(side=tk.LEFT,padx=10); e.focus_set(); e.bind("<Return>",lambda _: self._confirmar())
        self.lbl_error=tk.Label(self,text="",bg="#F0F2F5",fg="#E74C3C",font=("Segoe UI",9)); self.lbl_error.pack()
        bf=tk.Frame(self,bg="#F0F2F5"); bf.pack(fill=tk.X,padx=20,pady=(4,12))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Confirmar  ",command=self._confirmar).pack(side=tk.RIGHT)

    def _confirmar(self):
        nuevo=self.sv_cod.get().strip()
        if not nuevo: messagebox.showwarning("Atención","Ingresá un código.",parent=self); return
        if nuevo.upper() in CODIGOS_GENERICOS: messagebox.showwarning("Atención","No podés usar CH o MOP.",parent=self); return
        with get_conn() as conn:
            if conn.execute("SELECT 1 FROM tareas WHERE cod_tar=? AND id!=? LIMIT 1",(nuevo,self.tar_id)).fetchone():
                messagebox.showwarning("Atención",f"El código \"{nuevo}\" ya existe.",parent=self); return
            conn.execute("UPDATE tareas SET cod_tar=? WHERE id=?",(nuevo,self.tar_id))
        cb=self.on_asignado; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(nuevo))


# ═══════════════════════════════════════════════════════════
class InsumoDialog(tk.Toplevel):
    def __init__(self,parent,mode,cod_tar,on_save,cod_ins=None,cantidad_actual="1.0"):
        super().__init__(parent)
        self.mode=mode; self.cod_tar=cod_tar; self.cod_ins=cod_ins; self.on_save=on_save
        with get_conn() as conn:
            row=conn.execute("SELECT descripcion FROM tareas WHERE cod_tar=? AND activo=1 LIMIT 1",(cod_tar,)).fetchone()
        desc_tar=(row[0] or "")[:55] if row else ""
        titulo="Agregar insumo" if mode=="add" else "Editar insumo"
        self.title(f"{titulo}  —  {cod_tar}"); self.configure(bg="#F0F2F5"); self.resizable(True,True); self.grab_set()
        if mode=="add": self.geometry("760x520"); self.minsize(600,420)
        else: self.geometry("480x220"); self.minsize(400,200); self.resizable(False,False)
        self._build_ui(cod_tar,desc_tar,cantidad_actual)
        if mode=="add": self._load_insumos()

    def _build_ui(self,cod_tar,desc_tar,cantidad_actual):
        hdr=tk.Frame(self,bg="#2C3E50",height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        titulo="Agregar insumo" if self.mode=="add" else "Editar cantidad"
        tk.Label(hdr,text=f"  {titulo}  ·  {cod_tar} — {desc_tar}",bg="#2C3E50",fg="white",font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        if self.mode=="add":
            ff=tk.Frame(self,bg="#F0F2F5"); ff.pack(fill=tk.X,padx=12,pady=(8,4))
            ttk.Label(ff,text="Buscar insumo:").pack(side=tk.LEFT)
            self.sv_search=tk.StringVar(); self.sv_search.trace_add("write",lambda *_: self._filter())
            ttk.Entry(ff,textvariable=self.sv_search).pack(side=tk.LEFT,padx=6,fill=tk.X,expand=True)
            self.btn_rename=ttk.Button(ff,text="✎ Renombrar",style="Edit.TButton",state=tk.DISABLED,command=self._open_renombrar); self.btn_rename.pack(side=tk.LEFT,padx=(4,0))
            self.btn_nuevo=ttk.Button(ff,text="＋ Nuevo insumo",style="Add.TButton",command=self._open_nuevo_insumo); self.btn_nuevo.pack(side=tk.LEFT,padx=(4,0))
            frm=tk.Frame(self,bg="#F0F2F5"); frm.pack(fill=tk.BOTH,expand=True,padx=12,pady=(0,4))
            self.tv=ttk.Treeview(frm,columns=("cod","desc","unidad"),show="headings",selectmode="browse")
            self.tv.heading("cod",text="Código"); self.tv.heading("desc",text="Descripción"); self.tv.heading("unidad",text="Unidad")
            self.tv.column("cod",width=85,anchor=tk.CENTER,stretch=False); self.tv.column("desc",width=450,minwidth=200); self.tv.column("unidad",width=70,anchor=tk.CENTER,stretch=False)
            sb=ttk.Scrollbar(frm,orient=tk.VERTICAL,command=self.tv.yview); self.tv.configure(yscrollcommand=sb.set)
            self.tv.pack(side=tk.LEFT,fill=tk.BOTH,expand=True); sb.pack(side=tk.LEFT,fill=tk.Y)
            self.tv.bind("<<TreeviewSelect>>",self._on_ins_list_select)
        else:
            info_frm=tk.Frame(self,bg="#F0F2F5"); info_frm.pack(fill=tk.X,padx=20,pady=(16,8))
            with get_conn() as conn:
                row=conn.execute("SELECT descripcion,unidad FROM insumos WHERE cod_ins=?",(self.cod_ins,)).fetchone()
            ttk.Label(info_frm,text="Insumo:",font=("Segoe UI",10,"bold")).grid(row=0,column=0,sticky=tk.W,pady=2)
            ttk.Label(info_frm,text=f"{self.cod_ins}  —  {row[0] if row else ''}  [{row[1] if row else ''}]").grid(row=0,column=1,sticky=tk.W,padx=8,pady=2)
        bot=tk.Frame(self,bg="#F0F2F5"); bot.pack(fill=tk.X,padx=12,pady=(4,10))
        ttk.Label(bot,text="Cantidad por unidad de tarea:").pack(side=tk.LEFT)
        self.sv_qty=tk.StringVar(value=str(cantidad_actual))
        qty_entry=ttk.Entry(bot,textvariable=self.sv_qty,width=12)
        qty_entry.pack(side=tk.LEFT,padx=8); qty_entry.focus_set(); qty_entry.select_range(0,tk.END)
        qty_entry.bind("<Return>",lambda _: self._guardar())
        ttk.Button(bot,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bot,text="  Guardar  ",command=self._guardar).pack(side=tk.RIGHT)

    def _load_insumos(self):
        with get_conn() as conn:
            self._all=conn.execute("SELECT cod_ins,descripcion,unidad FROM insumos WHERE activo=1 ORDER BY cod_ins").fetchall()
        self._populate(self._all)

    def _populate(self,rows):
        self.tv.delete(*self.tv.get_children())
        for r in rows: self.tv.insert("",tk.END,values=r)

    def _filter(self):
        q=self.sv_search.get().lower()
        self._populate(self._all if not q else [r for r in self._all if q in r[0].lower() or q in (r[1] or "").lower()])

    def _open_nuevo_insumo(self): NuevoInsumoDialog(self,on_creado=self._on_nuevo_creado)

    def _on_nuevo_creado(self,cod_ins,desc,unidad):
        self._load_insumos()
        for item in self.tv.get_children():
            if self.tv.item(item,"values")[0]==cod_ins: self.tv.selection_set(item); self.tv.see(item); break

    def _on_ins_list_select(self,_=None):
        self.btn_rename.config(state=tk.NORMAL if self.tv.selection() else tk.DISABLED)

    def _open_renombrar(self):
        sel=self.tv.selection()
        if not sel: return
        vals=self.tv.item(sel[0],"values")
        RenombrarInsumoDialog(self,cod_ins=vals[0],desc_actual=vals[1],on_renombrado=self._on_renombrado)

    def _on_renombrado(self,cod_ins,nueva_desc):
        self._all=[(c,nueva_desc if c==cod_ins else d,u) for c,d,u in self._all]; self._filter()

    def _guardar(self):
        try:
            cantidad=float(self.sv_qty.get().replace(",",".")); 
            if cantidad<=0: raise ValueError
        except ValueError:
            messagebox.showwarning("Atención","Ingresá una cantidad válida (número positivo).",parent=self); return
        if self.mode=="add":
            sel=self.tv.selection()
            if not sel: messagebox.showwarning("Atención","Seleccioná un insumo de la lista.",parent=self); return
            vals=self.tv.item(sel[0],"values"); cod_ins=vals[0]; unidad=vals[2]
        else:
            cod_ins=self.cod_ins
            with get_conn() as conn:
                row=conn.execute("SELECT unidad FROM insumos WHERE cod_ins=?",(cod_ins,)).fetchone()
            unidad=row[0] if row else ""
        with get_conn() as conn:
            existente=conn.execute("SELECT i.cod_ins,ins.descripcion,i.cantidad FROM incidencias i LEFT JOIN insumos ins ON ins.cod_ins=i.cod_ins WHERE i.cod_tar=? AND i.cod_ins=?",(self.cod_tar,cod_ins)).fetchone()
        if existente:
            with get_conn() as conn:
                todos=conn.execute("SELECT i.cod_ins,ins.descripcion,i.cantidad FROM incidencias i LEFT JOIN insumos ins ON ins.cod_ins=i.cod_ins WHERE i.cod_tar=? ORDER BY i.cod_ins",(self.cod_tar,)).fetchall()
            idx=next((j for j,r in enumerate(todos) if r[0]==cod_ins),None)
            vecinos=[]
            if idx is not None:
                if idx>0: vecinos.append(todos[idx-1])
                vecinos.append(todos[idx])
                if idx<len(todos)-1: vecinos.append(todos[idx+1])
            lineas=["Va a reemplazar este insumo:\n"]
            for c,d,q in vecinos:
                marca="  ►" if c==cod_ins else "   "
                lineas.append(f"{marca} {c}  {(d or '')[:40]}  [{q}]")
            lineas.append(f"\n  Nuevo valor: cantidad = {cantidad}")
            if not messagebox.askyesno("Confirmar reemplazo","\n".join(lineas),parent=self): return
        with get_conn() as conn:
            conn.execute("""INSERT INTO incidencias (cod_tar,cod_ins,unidad,cantidad) VALUES (?,?,?,?)
                ON CONFLICT(cod_tar,cod_ins) DO UPDATE SET unidad=excluded.unidad,cantidad=excluded.cantidad""",
                (self.cod_tar,cod_ins,unidad,cantidad))
        accion="agregado" if self.mode=="add" else "actualizado"
        messagebox.showinfo("Guardado",f"Insumo  {cod_ins}  {accion}  ·  cantidad: {cantidad:.4f}",parent=self)
        self.on_save(); self.destroy()


# ═══════════════════════════════════════════════════════════
class RenombrarInsumoDialog(tk.Toplevel):
    def __init__(self,parent,cod_ins,desc_actual,on_renombrado):
        super().__init__(parent)
        self.cod_ins=cod_ins; self.on_renombrado=on_renombrado
        self.title(f"Renombrar insumo — {cod_ins}"); self.geometry("500x210")
        self.resizable(False,False); self.configure(bg="#F0F2F5"); self.grab_set()
        hdr=tk.Frame(self,bg="#2980B9",height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr,text=f"  Renombrar insumo  ·  {cod_ins}",bg="#2980B9",fg="white",font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        frm_actual=tk.Frame(self,bg="#F0F2F5"); frm_actual.pack(fill=tk.X,padx=20,pady=(14,4))
        ttk.Label(frm_actual,text="Nombre actual:",font=("Segoe UI",10,"bold")).grid(row=0,column=0,sticky=tk.W)
        ttk.Label(frm_actual,text=desc_actual,foreground="#555").grid(row=0,column=1,sticky=tk.W,padx=10)
        frm_nuevo=tk.Frame(self,bg="#F0F2F5"); frm_nuevo.pack(fill=tk.X,padx=20,pady=(0,4))
        ttk.Label(frm_nuevo,text="Nuevo nombre:").grid(row=0,column=0,sticky=tk.W)
        self.sv_desc=tk.StringVar(value=desc_actual)
        entry=ttk.Entry(frm_nuevo,textvariable=self.sv_desc,width=38,font=("Segoe UI",10))
        entry.grid(row=0,column=1,sticky=tk.W,padx=10); entry.focus_set(); entry.select_range(0,tk.END)
        entry.bind("<Return>",lambda _: self._confirmar())
        self.lbl_error=tk.Label(self,text="",bg="#F0F2F5",fg="#E74C3C",font=("Segoe UI",9)); self.lbl_error.pack()
        bf=tk.Frame(self,bg="#F0F2F5"); bf.pack(fill=tk.X,padx=20,pady=(2,12))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Guardar  ",command=self._confirmar).pack(side=tk.RIGHT)

    def _confirmar(self):
        nueva_desc=self.sv_desc.get().strip()
        if not nueva_desc: self.lbl_error.config(text="El nombre no puede estar vacío."); return
        with get_conn() as conn:
            conn.execute("UPDATE insumos SET descripcion=? WHERE cod_ins=?",(nueva_desc,self.cod_ins))
            cols=[r[1] for r in conn.execute("PRAGMA table_info(incidencias)").fetchall()]
            if "descripcion" in cols:
                conn.execute("UPDATE incidencias SET descripcion=? WHERE cod_ins=?",(nueva_desc,self.cod_ins))
        cb=self.on_renombrado; cod=self.cod_ins; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(cod,nueva_desc))


# ═══════════════════════════════════════════════════════════
class CambiarCodigoDialog(tk.Toplevel):
    def __init__(self,parent,tipo,cod_actual,on_cambiado):
        super().__init__(parent)
        self.tipo=tipo; self.cod_actual=cod_actual; self.on_cambiado=on_cambiado
        self.tabla="tareas" if tipo=="TAR" else "insumos"
        self.col_cod="cod_tar" if tipo=="TAR" else "cod_ins"
        self.col_inc="cod_tar" if tipo=="TAR" else "cod_ins"
        self.title(f"Nuevo COD_{tipo}  —  {cod_actual}"); self.geometry("460x220")
        self.resizable(False,False); self.configure(bg="#F0F2F5"); self.grab_set()
        color_hdr="#8E44AD" if tipo=="TAR" else "#2980B9"
        hdr=tk.Frame(self,bg=color_hdr,height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr,text=f"  Nuevo COD_{tipo}  ·  {cod_actual}",bg=color_hdr,fg="white",font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        campos=tk.Frame(self,bg="#F0F2F5"); campos.pack(fill=tk.X,padx=20,pady=(16,6))
        ttk.Label(campos,text="Código actual:",font=("Segoe UI",10,"bold")).grid(row=0,column=0,sticky=tk.W,pady=4)
        ttk.Label(campos,text=self.cod_actual,foreground="#555").grid(row=0,column=1,sticky=tk.W,padx=10,pady=4)
        ttk.Label(campos,text="Nuevo código:").grid(row=1,column=0,sticky=tk.W,pady=4)
        self.sv_nuevo=tk.StringVar()
        entry=ttk.Entry(campos,textvariable=self.sv_nuevo,width=28,font=("Segoe UI",10))
        entry.grid(row=1,column=1,sticky=tk.W,padx=10,pady=4); entry.focus_set()
        entry.bind("<Return>",lambda _: self._confirmar())
        aviso=tk.Frame(self,bg="#FEF9E7",bd=1,relief=tk.SOLID); aviso.pack(fill=tk.X,padx=20,pady=(0,4))
        tk.Label(aviso,text="  Se actualizará automáticamente en incidencias  ",bg="#FEF9E7",fg="#7D6608",font=("Segoe UI",9)).pack(pady=3)
        self.lbl_error=tk.Label(self,text="",bg="#F0F2F5",fg="#E74C3C",font=("Segoe UI",9)); self.lbl_error.pack()
        bf=tk.Frame(self,bg="#F0F2F5"); bf.pack(fill=tk.X,padx=20,pady=(2,12))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Guardar  ",command=self._confirmar).pack(side=tk.RIGHT)

    def _confirmar(self):
        nuevo=self.sv_nuevo.get().strip()
        if not nuevo: self.lbl_error.config(text="El nuevo código no puede estar vacío."); return
        if nuevo==self.cod_actual: self.lbl_error.config(text="El nuevo código es igual al actual."); return
        with get_conn() as conn:
            if conn.execute(f"SELECT 1 FROM {self.tabla} WHERE {self.col_cod}=? LIMIT 1",(nuevo,)).fetchone():
                self.lbl_error.config(text=f"El código «{nuevo}» ya existe."); return
        try:
            conn=get_conn()
            conn.execute(f"UPDATE {self.tabla} SET {self.col_cod}=? WHERE {self.col_cod}=?",(nuevo,self.cod_actual))
            conn.execute(f"UPDATE OR REPLACE incidencias SET {self.col_inc}=? WHERE {self.col_inc}=?",(nuevo,self.cod_actual))
            conn.commit(); conn.close()
        except Exception as e:
            try: conn.rollback(); conn.close()
            except: pass
            self.lbl_error.config(text=f"Error: {e}"); return
        cb=self.on_cambiado; viejo=self.cod_actual; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(viejo,nuevo))


# ═══════════════════════════════════════════════════════════
class NuevoInsumoDialog(tk.Toplevel):
    def __init__(self,parent,on_creado):
        super().__init__(parent)
        self.on_creado=on_creado; self.title("Nuevo insumo"); self.geometry("460x240")
        self.resizable(False,False); self.configure(bg="#F0F2F5"); self.grab_set()
        hdr=tk.Frame(self,bg="#27AE60",height=44); hdr.pack(fill=tk.X); hdr.pack_propagate(False)
        tk.Label(hdr,text="  Nuevo insumo",bg="#27AE60",fg="white",font=("Segoe UI",11,"bold")).pack(side=tk.LEFT,padx=8,pady=8)
        campos=tk.Frame(self,bg="#F0F2F5"); campos.pack(fill=tk.X,padx=20,pady=(16,6))
        ttk.Label(campos,text="Código (COD_INS):").grid(row=0,column=0,sticky=tk.W,pady=4)
        self.sv_cod=tk.StringVar()
        self.entry_cod=ttk.Entry(campos,textvariable=self.sv_cod,width=20,font=("Segoe UI",10))
        self.entry_cod.grid(row=0,column=1,sticky=tk.W,padx=10,pady=4); self.entry_cod.focus_set()
        ttk.Label(campos,text="Descripción:").grid(row=1,column=0,sticky=tk.W,pady=4)
        self.sv_desc=tk.StringVar()
        ttk.Entry(campos,textvariable=self.sv_desc,width=32,font=("Segoe UI",10)).grid(row=1,column=1,sticky=tk.W,padx=10,pady=4)
        ttk.Label(campos,text="Unidad:").grid(row=2,column=0,sticky=tk.W,pady=4)
        self.sv_unid=tk.StringVar()
        ttk.Entry(campos,textvariable=self.sv_unid,width=12,font=("Segoe UI",10)).grid(row=2,column=1,sticky=tk.W,padx=10,pady=4)
        self.lbl_error=tk.Label(self,text="",bg="#F0F2F5",fg="#E74C3C",font=("Segoe UI",9)); self.lbl_error.pack()
        bf=tk.Frame(self,bg="#F0F2F5"); bf.pack(fill=tk.X,padx=20,pady=(4,12))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Guardar  ",command=self._guardar).pack(side=tk.RIGHT)
        self.entry_cod.bind("<Return>",lambda _: self.focus_next())
        self.bind("<Return>",lambda _: self._guardar())

    def focus_next(self): self.tk_focusNext().focus_set()

    def _guardar(self):
        cod=self.sv_cod.get().strip(); desc=self.sv_desc.get().strip(); unid=self.sv_unid.get().strip()
        if not cod: self.lbl_error.config(text="El código no puede estar vacío."); return
        if not desc: self.lbl_error.config(text="La descripción no puede estar vacía."); return
        with get_conn() as conn:
            if conn.execute("SELECT 1 FROM insumos WHERE cod_ins=? LIMIT 1",(cod,)).fetchone():
                self.lbl_error.config(text=f"El código «{cod}» ya existe en insumos."); return
            conn.execute("INSERT INTO insumos (cod_ins,descripcion,unidad,activo) VALUES (?,?,?,1)",(cod,desc,unid))
        cb=self.on_creado; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(cod,desc,unid))


# ═══════════════════════════════════════════════════════════
class ModificarTareaDialog(tk.Toplevel):
    def __init__(self,master,tar_id,row,on_guardado):
        super().__init__(master)
        self.on_guardado=on_guardado; self.tar_id=tar_id
        self.cod_original,desc_orig,unid_orig,self.id_rs_orig=row
        self.title(f"Modificar tarea  ·  {self.cod_original}")
        self.resizable(False,False); self.grab_set(); self.configure(bg="#F0F2F5")
        with get_conn() as conn:
            tiene_rubros=conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='rubros'").fetchone()[0]
            if tiene_rubros:
                rubs_raw=conn.execute("SELECT cod_rubro,designacion,id_rs_ini FROM rubros WHERE es_padre=1 ORDER BY id_rs_ini").fetchall()
                subs_raw=conn.execute("""SELECT s.cod_rubro,s.id_rs_ini,s.designacion,COALESCE(s.cod_padre,s.cod_rubro) AS rub_padre FROM rubros s
                    WHERE s.es_padre=0 OR NOT EXISTS (SELECT 1 FROM rubros c WHERE c.cod_padre=s.cod_rubro)
                    ORDER BY COALESCE((SELECT p.id_rs_ini FROM rubros p WHERE p.cod_rubro=s.cod_padre),s.id_rs_ini),s.cod_rubro""").fetchall()
            else:
                rubs_raw=[]; subs_raw=[]
            cols_tareas=[r[1] for r in conn.execute("PRAGMA table_info(tareas)").fetchall()]
            qty_row=conn.execute("SELECT cantidad FROM tareas WHERE id=?",(tar_id,)).fetchone() if "cantidad" in cols_tareas else None
            cantidad_orig=qty_row[0] if qty_row and qty_row[0] is not None else ""
        self._rubs=[(r[0],f"{r[0]}  —  {r[1]}") for r in rubs_raw]
        self._subs_por_rub={}; self._sub_info={}
        for s in subs_raw:
            cod_s,id_rs,des_s,padre=s; label_s=f"{cod_s}  —  {des_s}"
            self._subs_por_rub.setdefault(padre,[]).append(label_s); self._sub_info[label_s]=(cod_s,id_rs)
        cur_rub_label=""; cur_sub_label=""
        if self.id_rs_orig and subs_raw:
            sub_match=next((s for s in subs_raw if s[1]==self.id_rs_orig),None)
            if sub_match:
                cur_sub_label=f"{sub_match[0]}  —  {sub_match[2]}"
                rub_match=next((r for r in rubs_raw if r[0]==sub_match[3]),None)
                if rub_match: cur_rub_label=f"{rub_match[0]}  —  {rub_match[1]}"
        self.sv_rub=tk.StringVar(value=cur_rub_label); self.sv_sub=tk.StringVar(value=cur_sub_label)
        self.sv_cod=tk.StringVar(value=self.cod_original); self.sv_desc=tk.StringVar(value=desc_orig or "")
        self.sv_unid=tk.StringVar(value=unid_orig or ""); self.sv_cantidad=tk.StringVar(value=str(cantidad_orig))
        hdr=tk.Frame(self,bg="#1F5C99"); hdr.pack(fill=tk.X)
        tk.Label(hdr,text=f"  Modificar  ·  {self.cod_original}",bg="#1F5C99",fg="white",font=("Segoe UI",11,"bold")).pack(anchor=tk.W,padx=10,pady=8)
        frm=ttk.Frame(self,padding=16); frm.pack(fill=tk.BOTH)
        for i,lbl in enumerate(["Rubro:","SubRubro:","COD_TAR:","Descripción:","Unidad:","Cantidad:"]):
            ttk.Label(frm,text=lbl).grid(row=i,column=0,sticky=tk.W,pady=4)
        self.cb_rub=ttk.Combobox(frm,textvariable=self.sv_rub,values=[r[1] for r in self._rubs],width=46,state="readonly")
        self.cb_rub.grid(row=0,column=1,sticky=tk.W,padx=(8,0),pady=4)
        self.cb_sub=ttk.Combobox(frm,textvariable=self.sv_sub,values=self._subs_por_rub.get(next((r[0] for r in self._rubs if r[1]==cur_rub_label),""),[]),width=46,state="readonly")
        self.cb_sub.grid(row=1,column=1,sticky=tk.W,padx=(8,0),pady=4)
        def on_rub_select(_=None):
            sel=self.sv_rub.get(); cod_rub=next((r[0] for r in self._rubs if r[1]==sel),None)
            subs=self._subs_por_rub.get(cod_rub,[]); self.cb_sub.config(values=subs); self.sv_sub.set("")
            if subs: self.cb_sub.focus_set()
        self.cb_rub.bind("<<ComboboxSelected>>",on_rub_select)
        entries=[]
        for i,(sv,w) in enumerate([(self.sv_cod,18),(self.sv_desc,46),(self.sv_unid,10),(self.sv_cantidad,14)],start=2):
            e=ttk.Entry(frm,textvariable=sv,width=w,font=("Segoe UI",10)); e.grid(row=i,column=1,sticky=tk.W,padx=(8,0),pady=4); entries.append(e)
        self.cb_sub.bind("<Return>",lambda _: entries[0].focus_set())
        for i in range(len(entries)-1): entries[i].bind("<Return>",lambda _,j=i+1: entries[j].focus_set())
        entries[-1].bind("<Return>",lambda _: self._guardar())
        self.lbl_error=ttk.Label(frm,text="",foreground="red",font=("Segoe UI",9)); self.lbl_error.grid(row=6,column=0,columnspan=2,sticky=tk.W,pady=(4,0))
        bf=ttk.Frame(frm); bf.grid(row=7,column=0,columnspan=2,sticky=tk.E,pady=(12,0))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Guardar  ",command=self._guardar).pack(side=tk.RIGHT)
        if cur_rub_label: entries[0].focus_set(); entries[0].selection_range(0,tk.END)
        else: self.cb_rub.focus_set()
        self.update_idletasks(); self.geometry(f"+{master.winfo_rootx()+80}+{master.winfo_rooty()+80}")

    def _guardar(self):
        sub_label=self.sv_sub.get().strip(); cod=self.sv_cod.get().strip().upper()
        desc=self.sv_desc.get().strip(); unid=self.sv_unid.get().strip(); qty_str=self.sv_cantidad.get().strip()
        if not sub_label: self.lbl_error.config(text="Seleccioná un SubRubro."); return
        if not cod: self.lbl_error.config(text="El código no puede estar vacío."); return
        if not desc: self.lbl_error.config(text="La descripción no puede estar vacía."); return
        cantidad=None
        if qty_str:
            try:
                cantidad=float(qty_str.replace(",",".")); 
                if cantidad<=0: raise ValueError
            except ValueError: self.lbl_error.config(text="Cantidad inválida."); return
        sub_match=self._sub_info.get(sub_label); id_rs=sub_match[1] if sub_match else self.id_rs_orig
        if id_rs and id_rs!=self.id_rs_orig:
            with get_conn() as conn:
                usado=conn.execute("SELECT id_tarea FROM tareas WHERE id_rs=? AND id_tarea IS NOT NULL AND id!=? ORDER BY id_tarea",(id_rs,self.tar_id)).fetchall()
            CHARS="0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
            def to_b36(n,w):
                if n==0: return "0"*w
                s=""
                while n: s=CHARS[n%36]+s; n//=36
                return s.zfill(w)
            def from_b36(s): return sum(CHARS.index(c)*(36**i) for i,c in enumerate(reversed(s.upper())))
            usados_num={from_b36(r[0][2:]) for r in usado if r[0]}
            seq=0
            while seq in usados_num: seq+=1
            id_tarea=id_rs+to_b36(seq,2)
        else:
            with get_conn() as conn:
                r=conn.execute("SELECT id_tarea FROM tareas WHERE id=?",(self.tar_id,)).fetchone()
            id_tarea=r[0] if r else None
        try:
            with get_conn() as conn:
                if cod!=self.cod_original:
                    if conn.execute("SELECT 1 FROM tareas WHERE cod_tar=?",(cod,)).fetchone():
                        self.lbl_error.config(text=f"El código «{cod}» ya existe."); return
                    conn.execute("UPDATE incidencias    SET cod_tar=? WHERE cod_tar=?",(cod,self.cod_original))
                    conn.execute("UPDATE incidencias_v2 SET cod_tar=? WHERE cod_tar=?",(cod,self.cod_original))
                cols=[r[1] for r in conn.execute("PRAGMA table_info(tareas)").fetchall()]
                if "cantidad" not in cols: conn.execute("ALTER TABLE tareas ADD COLUMN cantidad REAL")
                conn.execute("UPDATE tareas SET cod_tar=?,descripcion=?,unidad=?,id_rs=?,id_tarea=?,cantidad=? WHERE id=?",
                             (cod,desc,unid,id_rs,id_tarea,cantidad,self.tar_id))
        except Exception as e:
            self.lbl_error.config(text=f"Error al guardar: {e}"); return
        cod_viejo=self.cod_original; cb=self.on_guardado; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(cod_viejo,cod))


# ═══════════════════════════════════════════════════════════
class AgregarTareaDialog(tk.Toplevel):
    def __init__(self,master,on_guardado):
        super().__init__(master)
        self.on_guardado=on_guardado; self.title("Agregar tarea"); self.resizable(False,False); self.grab_set()
        hdr=tk.Frame(self,bg="#27AE60"); hdr.pack(fill=tk.X)
        tk.Label(hdr,text="  Nueva tarea",bg="#27AE60",fg="white",font=("Segoe UI",11,"bold")).pack(anchor=tk.W,padx=10,pady=8)
        frm=ttk.Frame(self,padding=16); frm.pack(fill=tk.BOTH)
        with get_conn() as conn:
            tiene_rubros=conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='rubros'").fetchone()[0]
            if tiene_rubros:
                rubs_raw=conn.execute("SELECT cod_rubro,designacion,id_rs_ini FROM rubros WHERE es_padre=1 ORDER BY id_rs_ini").fetchall()
                subs_raw=conn.execute("""SELECT s.cod_rubro,s.id_rs_ini,s.designacion,COALESCE(s.cod_padre,s.cod_rubro) AS rub_padre FROM rubros s
                    WHERE s.es_padre=0 OR NOT EXISTS (SELECT 1 FROM rubros c WHERE c.cod_padre=s.cod_rubro)
                    ORDER BY COALESCE((SELECT p.id_rs_ini FROM rubros p WHERE p.cod_rubro=s.cod_padre),s.id_rs_ini),s.cod_rubro""").fetchall()
            else:
                rubs_raw=[]; subs_raw=[]
        self._rubs=[(r[0],f"{r[0]}  —  {r[1]}") for r in rubs_raw]
        self._subs_por_rub={}; self._sub_info={}
        for s in subs_raw:
            cod_s,id_rs,des_s,padre=s; label_s=f"{cod_s}  —  {des_s}"
            self._subs_por_rub.setdefault(padre,[]).append(label_s); self._sub_info[label_s]=(cod_s,id_rs)
        self.sv_rub=tk.StringVar(); self.sv_sub=tk.StringVar()
        self.sv_cod=tk.StringVar(); self.sv_desc=tk.StringVar(); self.sv_unid=tk.StringVar()
        for i,lbl in enumerate(["Rubro","SubRubro","COD_TAR","Descripción","Unidad"]):
            ttk.Label(frm,text=lbl+":").grid(row=i,column=0,sticky=tk.W,pady=4)
        self.cb_rub=ttk.Combobox(frm,textvariable=self.sv_rub,values=[r[1] for r in self._rubs],width=46,state="readonly")
        self.cb_rub.grid(row=0,column=1,sticky=tk.W,padx=(8,0),pady=4)
        self.cb_sub=ttk.Combobox(frm,textvariable=self.sv_sub,values=[],width=46,state="readonly")
        self.cb_sub.grid(row=1,column=1,sticky=tk.W,padx=(8,0),pady=4)
        def on_rub_select(_=None):
            sel=self.sv_rub.get(); cod_rub=next((r[0] for r in self._rubs if r[1]==sel),None)
            subs=self._subs_por_rub.get(cod_rub,[]); self.cb_sub.config(values=subs); self.sv_sub.set("")
            if subs: self.cb_sub.focus_set()
        self.cb_rub.bind("<<ComboboxSelected>>",on_rub_select)
        entries=[]
        for i,(sv,w) in enumerate([(self.sv_cod,18),(self.sv_desc,46),(self.sv_unid,10)],2):
            e=ttk.Entry(frm,textvariable=sv,width=w,font=("Segoe UI",10)); e.grid(row=i,column=1,sticky=tk.W,padx=(8,0),pady=4); entries.append(e)
        self.cb_sub.bind("<Return>",lambda _: entries[0].focus_set())
        entries[0].bind("<Return>",lambda _: entries[1].focus_set())
        entries[1].bind("<Return>",lambda _: entries[2].focus_set())
        entries[2].bind("<Return>",lambda _: self._guardar())
        self.lbl_error=ttk.Label(frm,text="",foreground="red",font=("Segoe UI",9)); self.lbl_error.grid(row=5,column=0,columnspan=2,sticky=tk.W,pady=(4,0))
        bf=ttk.Frame(frm); bf.grid(row=6,column=0,columnspan=2,sticky=tk.E,pady=(12,0))
        ttk.Button(bf,text="Cancelar",command=self.destroy).pack(side=tk.RIGHT,padx=(4,0))
        ttk.Button(bf,text="  Guardar  ",command=self._guardar).pack(side=tk.RIGHT)
        self.cb_rub.focus_set(); self.update_idletasks()
        self.geometry(f"+{master.winfo_rootx()+80}+{master.winfo_rooty()+80}")

    def _guardar(self):
        sub_label=self.sv_sub.get().strip(); cod=self.sv_cod.get().strip().upper()
        desc=self.sv_desc.get().strip(); unid=self.sv_unid.get().strip()
        if not sub_label: self.lbl_error.config(text="Seleccioná un SubRubro."); return
        if not cod: self.lbl_error.config(text="El código no puede estar vacío."); return
        if not desc: self.lbl_error.config(text="La descripción no puede estar vacía."); return
        sub_match=self._sub_info.get(sub_label); id_rs=sub_match[1] if sub_match else None
        id_tarea=None
        if id_rs:
            with get_conn() as conn:
                usado=conn.execute("SELECT id_tarea FROM tareas WHERE id_rs=? AND id_tarea IS NOT NULL ORDER BY id_tarea",(id_rs,)).fetchall()
            CHARS='0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'
            def to_b36(n,w):
                if n==0: return '0'*w
                s=''
                while n: s=CHARS[n%36]+s; n//=36
                return s.zfill(w)
            def from_b36(s): return sum(CHARS.index(c)*(36**i) for i,c in enumerate(reversed(s.upper())))
            ss_usados={from_b36(r[0][2:]) for r in usado if r[0]}
            ss=0
            while ss in ss_usados: ss+=1
            id_tarea=id_rs+to_b36(ss,2)
        with get_conn() as conn:
            if conn.execute("SELECT 1 FROM tareas WHERE cod_tar=?",(cod,)).fetchone():
                self.lbl_error.config(text=f"El código «{cod}» ya existe."); return
            conn.execute("INSERT INTO tareas (cod_tar,descripcion,unidad,activo,id_tarea,id_rs) VALUES (?,?,?,1,?,?)",(cod,desc,unid,id_tarea,id_rs))
        cb=self.on_guardado; self.grab_release(); self.destroy()
        self.master.after(50,lambda: cb(cod))


# ═══════════════════════════════════════════════════════════
def main():
    """Punto de entrada — acepta --db y --title desde línea de comandos."""
    import argparse
    global DB_PATH
    parser = argparse.ArgumentParser(description="Gestión de Tareas e Insumos")
    parser.add_argument("--db",    default=DB_PATH,
                        help="Ruta a la base de datos SQLite")
    parser.add_argument("--title", default="Gestión de Tareas e Insumos",
                        help="Título de la ventana")
    args = parser.parse_args()
    DB_PATH = os.path.normpath(args.db)
    app = App()
    app.title(args.title)
    app.mainloop()


if __name__ == "__main__":
    main()
