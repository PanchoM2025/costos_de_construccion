"""
maestro.py
Lee Presupuestos.xlsx y PUnitarios_*.xlsx y expone los datos como estructuras Python.

Variantes de tarea:
  ' '  Mano de obra propia
  'S'  Subcontrato
  'E'  Subcontrato con H° Elaborado
  'X'  Subcontrato solo mano de obra
"""
import openpyxl
from pathlib import Path

SUFIJOS = {' ': 'M.O. Propia', 'S': 'Subcontrato', 'E': 'H° Elaborado', 'X': 'Solo M.O.'}


def _separar_variante(cod: str) -> tuple[str, str]:
    """Devuelve (codigo_base, sufijo). Sufijo es ' ' si no tiene letra."""
    if cod and cod[-1] in ('S', 'E', 'X'):
        return cod[:-1], cod[-1]
    return cod, ' '


def cargar_maestro(ruta_presupuestos: str) -> dict:
    """
    Lee Presupuestos.xlsx y devuelve:
    {
      'rubros': ['TAREAS PRELIMINARES', ...],
      'tareas': {
          'TAREAS PRELIMINARES': [
              {'cod': '0110', 'nombre': 'Limpieza superficial.', 'unidad': 'm2'},
              ...
          ]
      },
      'tareas_por_cod': {'0110': {...}, ...},
      'insumos': {'000114': {'nombre': 'Oficial.', 'unidad': 'hs'}, ...},
      'incidencias': {'0110': [{'cod_ins': '000114', 'unidad': 'hs', 'cantidad': 0.1}, ...]}
    }
    """
    wb = openpyxl.load_workbook(ruta_presupuestos, read_only=True, data_only=True)

    # ── Tareas y Rubros ──────────────────────────────────────────────────────
    rubros = []
    tareas_por_rubro = {}
    tareas_por_cod = {}
    rubro_actual = None

    ws_tareas = wb["Tareas"]
    for row in ws_tareas.iter_rows(min_row=2, values_only=True):
        cod, nombre, unidad = row[0], row[1], row[2]
        cod_str = str(cod).strip() if cod is not None else ""
        # Fila de rubro: código de 4 dígitos terminado en '000' (ej. 0000, 1000, 2000...)
        es_rubro = cod_str and len(cod_str) == 4 and cod_str.endswith("000")
        if es_rubro or (cod is None and nombre):
            rubro_actual = str(nombre).strip() if nombre else ""
            if rubro_actual:
                rubros.append(rubro_actual)
                tareas_por_rubro[rubro_actual] = []
        elif cod and rubro_actual:
            cod_str = str(cod).strip()
            base, variante = _separar_variante(cod_str)
            tarea = {
                "cod":      cod_str,
                "base":     base,
                "variante": variante,          # ' ', 'S', 'E' o 'X'
                "nombre":   str(nombre).strip() if nombre else "",
                "unidad":   str(unidad).strip() if unidad else "",
                "rubro":    rubro_actual,
            }
            tareas_por_rubro[rubro_actual].append(tarea)
            tareas_por_cod[cod_str] = tarea

    # ── Insumos ──────────────────────────────────────────────────────────────
    insumos = {}
    ws_ins = wb["Insumos"]
    for row in ws_ins.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        cod_ins = str(row[0]).strip()
        insumos[cod_ins] = {
            "cod": cod_ins,
            "nombre": str(row[1]).strip() if row[1] else "",
            "unidad": str(row[2]).strip() if row[2] else "",
        }

    # ── Incidencias ──────────────────────────────────────────────────────────
    incidencias = {}
    ws_inc = wb["Incidencias"]
    for row in ws_inc.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        cod_tar = str(row[0]).strip()
        cod_ins = str(row[1]).strip() if row[1] else None
        if not cod_ins:
            continue
        cant = row[3] if row[3] else 0
        if cod_tar not in incidencias:
            incidencias[cod_tar] = []
        incidencias[cod_tar].append({
            "cod_ins": cod_ins,
            "unidad": str(row[2]).strip() if row[2] else "",
            "cantidad": float(cant),
        })

    wb.close()

    return {
        "rubros": rubros,
        "tareas": tareas_por_rubro,
        "tareas_por_cod": tareas_por_cod,
        "insumos": insumos,
        "incidencias": incidencias,
    }


def cargar_precios(ruta_precios: str) -> dict:
    """
    Lee PUnitarios_*.xlsx y devuelve:
    {'000114': 1500.0, '000116': 1200.0, ...}
    Solo incluye los que tienen precio cargado (no None).
    """
    wb = openpyxl.load_workbook(ruta_precios, read_only=True, data_only=True)
    ws = wb["Insumos"]
    precios = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        cod = str(row[0]).strip()
        precio = row[3]
        if precio is not None:
            try:
                precios[cod] = float(precio)
            except (ValueError, TypeError):
                pass
    wb.close()
    return precios


def listar_archivos_precios(carpeta_data: str) -> list:
    """Devuelve lista de archivos .xlsx en la carpeta Data/Precios."""
    carpeta = Path(carpeta_data)
    if not carpeta.exists():
        return []
    return sorted(carpeta.glob("PUnitarios*.xlsx"), reverse=True)
