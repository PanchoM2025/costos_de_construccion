"""
generar_listado_para_pedido_precios.py
Genera un Excel virgen para pedido de precios con todos los insumos de Master_FINAL.db
"""

import sqlite3
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Instalando openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# ─── CONFIGURACIÓN ────────────────────────────────────────────────────────────
DB_PATH = r"C:\Dropbox\Analisis_de_costos_y_Presupuestacion\Costos\Data\Master_FINAL.db"
OUTPUT_PATH = r"C:\Dropbox\Analisis_de_costos_y_Presupuestacion\Costos\Data\Pedido_Precios.xlsx"
# ──────────────────────────────────────────────────────────────────────────────

def conectar(db_path):
    if not Path(db_path).exists():
        raise FileNotFoundError(f"No se encontró la base de datos en:\n  {db_path}")
    return sqlite3.connect(db_path)

def obtener_insumos(conn):
    cur = conn.cursor()
    # Detectar columnas disponibles
    cur.execute("PRAGMA table_info(insumos)")
    cols = {row[1].lower(): row[1] for row in cur.fetchall()}
    print(f"Columnas de la tabla insumos: {list(cols.keys())}")

    # Construir SELECT dinámico según columnas existentes
    col_codigo  = cols.get("cod_ins") or cols.get("codigo") or cols.get("id_insumo") or "rowid"
    col_nombre  = cols.get("descripcion") or cols.get("nombre") or cols.get("desc_ins") or cols.get("name") or "''"
    col_unidad  = cols.get("unidad") or cols.get("ud") or cols.get("um") or "''"
    col_tipo    = cols.get("tipo") or cols.get("categoria") or cols.get("rubro") or None

    select_cols = [col_codigo, col_nombre, col_unidad]
    if col_tipo:
        select_cols.append(col_tipo)

    query = f"""
        SELECT {', '.join(select_cols)}
        FROM insumos
        ORDER BY {col_codigo}
    """
    cur.execute(query)
    rows = cur.fetchall()
    return rows, col_tipo is not None

def crear_excel(insumos, tiene_tipo, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedido de Precios"

    # ── Estilos ──────────────────────────────────────────────────────────────
    AZUL        = "1F4E79"
    AZUL_CLARO  = "BDD7EE"
    GRIS        = "F2F2F2"
    NARANJA     = "FF6600"

    header_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    header_fill   = PatternFill("solid", fgColor=AZUL)
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_font    = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    title_fill    = PatternFill("solid", fgColor=AZUL)
    title_align   = Alignment(horizontal="center", vertical="center")

    data_font     = Font(name="Calibri", size=9)
    alt_fill      = PatternFill("solid", fgColor=GRIS)
    precio_fill   = PatternFill("solid", fgColor="FFFFD6")   # amarillo suave

    thin  = Side(style="thin",   color="BFBFBF")
    thick = Side(style="medium", color=AZUL)
    border_data   = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_precio = Border(left=thick, right=thick, top=thin, bottom=thin)

    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")

    # ── Título ───────────────────────────────────────────────────────────────
    num_cols = 6 if tiene_tipo else 5
    ws.merge_cells(f"A1:{get_column_letter(num_cols)}1")
    ws["A1"] = "PEDIDO DE PRECIOS DE INSUMOS"
    ws["A1"].font    = title_font
    ws["A1"].fill   = title_fill
    ws["A1"].alignment = title_align
    ws.row_dimensions[1].height = 24

    # Fecha
    from datetime import date
    ws.merge_cells(f"A2:{get_column_letter(num_cols)}2")
    ws["A2"] = f"Fecha: {date.today().strftime('%d/%m/%Y')}    Proveedor: ___________________________"
    ws["A2"].font      = Font(name="Calibri", italic=True, size=9)
    ws["A2"].alignment = left
    ws.row_dimensions[2].height = 16

    # ── Encabezados ──────────────────────────────────────────────────────────
    col = 1
    headers = []
    if tiene_tipo:
        headers.append(("TIPO", 12))
    headers += [
        ("CÓDIGO",       10),
        ("DESCRIPCIÓN",  55),
        ("UNIDAD",        9),
        ("PRECIO UNIT.", 14),
        ("OBSERVACIONES",30),
    ]

    for i, (h, w) in enumerate(headers):
        c = ws.cell(row=3, column=col + i, value=h)
        c.font      = header_font
        c.fill      = header_fill
        c.alignment = header_align
        c.border    = border_data
        ws.column_dimensions[get_column_letter(col + i)].width = w

    ws.row_dimensions[3].height = 28

    # ── Datos ────────────────────────────────────────────────────────────────
    for i, row in enumerate(insumos):
        r = 4 + i
        fill = alt_fill if i % 2 == 0 else PatternFill()

        col_idx = 1
        if tiene_tipo:
            c = ws.cell(row=r, column=col_idx, value=row[3] if len(row) > 3 else "")
            c.font = data_font; c.fill = fill; c.alignment = center; c.border = border_data
            col_idx += 1

        # Código
        c = ws.cell(row=r, column=col_idx, value=row[0])
        c.font = data_font; c.fill = fill; c.alignment = center; c.border = border_data
        col_idx += 1

        # Descripción
        c = ws.cell(row=r, column=col_idx, value=row[1])
        c.font = data_font; c.fill = fill; c.alignment = left; c.border = border_data
        col_idx += 1

        # Unidad
        c = ws.cell(row=r, column=col_idx, value=row[2])
        c.font = data_font; c.fill = fill; c.alignment = center; c.border = border_data
        col_idx += 1

        # Precio (en blanco, destacado)
        c = ws.cell(row=r, column=col_idx, value=None)
        c.font = Font(name="Calibri", size=9, bold=True)
        c.fill = precio_fill; c.alignment = center; c.border = border_precio
        c.number_format = '"$"#,##0.00'
        col_idx += 1

        # Observaciones (en blanco)
        c = ws.cell(row=r, column=col_idx, value=None)
        c.font = data_font; c.fill = fill; c.alignment = left; c.border = border_data

        ws.row_dimensions[r].height = 14

    # ── Freeze y filtros ─────────────────────────────────────────────────────
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(num_cols)}{3 + len(insumos)}"

    # ── Guardar ──────────────────────────────────────────────────────────────
    wb.save(output_path)
    print(f"\n✓ Excel generado con {len(insumos)} insumos:")
    print(f"  {output_path}")

# ─── MAIN ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print(f"Conectando a: {DB_PATH}")
    try:
        conn = conectar(DB_PATH)
        insumos, tiene_tipo = obtener_insumos(conn)
        conn.close()
        print(f"Insumos encontrados: {len(insumos)}")
        crear_excel(insumos, tiene_tipo, OUTPUT_PATH)
    except FileNotFoundError as e:
        print(f"\nERROR: {e}")
        sys.exit(1)
    except Exception as e:
        import traceback
        print(f"\nERROR inesperado: {e}")
        traceback.print_exc()
        sys.exit(1)
